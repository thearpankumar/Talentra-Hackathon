import re
import tempfile
from pathlib import Path

from ingestion_engine.config import parse_config
from ingestion_engine.generate import RunContext, gen_ids, generate_batch, to_copy_bytes
from ingestion_engine.runner import Runner

from .conftest import make_cfg

NULL = "\\N"


def test_introspection(settings):
    sch = Runner(settings).get_schema()
    order = sch.order()
    assert order.index("users") < order.index("candidate_profile") < order.index("application")
    assert order.index("country") < order.index("city")
    cp = sch.tables["candidate_profile"]
    assert cp.pk == ["id"] and ["user_id"] in cp.uniques
    assert cp.columns["gender"].enum_values and "undisclosed" in cp.columns["gender"].enum_values
    assert sch.tables["candidate_skill"].pk == ["candidate_profile_id", "skill_id"]
    assert not any(t.startswith("talentra_ingest_") for t in sch.tables)
    assert cp.fk_for("user_id").ref_table == "users"


def _ctx(settings, bs=500):
    r = Runner(settings)
    return r, RunContext(r.get_schema(), "t", Path(tempfile.mkdtemp()), bs, conn_factory=r.connect)


def test_generation_is_deterministic_and_respects_rules(settings):
    r, ctx = _ctx(settings)
    sc = parse_config(make_cfg(users=1000)).scenarios[0]
    ts = next(t for t in sc.tables if t.table == "users")
    a1, cols1, st = generate_batch(ctx, sc, ts, 0, 500, 1000)
    a2, cols2, _ = generate_batch(ctx, sc, ts, 0, 500, 1000)
    assert cols1 == cols2 and all((a1[c] == a2[c]).all() for c in cols1)
    b1, _, _ = generate_batch(ctx, sc, ts, 1, 500, 1000)
    assert not set(a1["id"]) & set(b1["id"])
    assert all(e == e.lower() for e in a1["email"])  # lowercase emails
    assert all(re.fullmatch(r"\+[1-9]\d{6,14}", p) or p == NULL for p in a1["phone_e164"])  # E.164
    assert len(set(a1["email"])) == 500
    assert "created_by" not in cols1  # DB default applies
    names = "".join(a1["full_name"])
    assert re.search("[ऀ-ॿ]", names) and re.search("[؀-ۿ]", names) and re.search("[一-鿿]", names)
    assert st["dups"] > 0
    assert len(to_copy_bytes(a1, cols1).splitlines()) == 500


def test_unknown_vs_not_applicable_and_pair_checks(settings):
    r, ctx = _ctx(settings)
    d = make_cfg(users=3000)
    d["scenarios"][0]["null_rate"] = 0.3
    sc = parse_config(d).scenarios[0]
    ts = next(t for t in sc.tables if t.table == "candidate_profile")
    a, cols, st = generate_batch(ctx, sc, ts, 0, 500, 5000)
    disc = a["salary_disclosure"]
    assert {"unknown", "not_applicable", "disclosed"} <= set(disc)
    for i in range(500):
        null = a["expected_salary_min"][i] == NULL
        assert null == (disc[i] != "disclosed") == (a["expected_salary_max"][i] == NULL)
        if not null:
            assert float(a["expected_salary_min"][i]) <= float(a["expected_salary_max"][i])
    assert set(a["gender"]) - {NULL} <= {"female", "male", "other", "undisclosed"}


def test_ids_deterministic_and_unique():
    a = gen_ids("uuid", 1, "t", 0, 1000)
    assert (a == gen_ids("uuid", 1, "t", 0, 1000)).all() and len(set(a)) == 1000
    assert not set(a) & set(gen_ids("uuid", 2, "t", 0, 1000))


def test_each_scenario_has_its_own_configured_distribution(settings):
    """Two scenarios with different distribution specs yield statistically different (and spec-matching) data."""
    import numpy as np
    r, ctx = _ctx(settings, bs=4000)
    d = make_cfg(users=4000, extra_tables=False)
    s1 = d["scenarios"][0]
    s1["tables"][1]["fields"] = {"years_experience": {"dist": "normal", "params": {"mean": 8, "std": 2}, "min": 0, "max": 40, "round": 1},
                                 "expected_salary_min": {"dist": "lognormal", "params": {"mu": 13.0, "sigma": 0.3}}}
    s2 = {**s1, "id": "other", "seed": 5, "locale_mix": {"ar_AE": 1.0}}
    s2["tables"] = [dict(s1["tables"][0]), {**s1["tables"][1], "fields": {"years_experience": {"dist": "exponential", "params": {"scale": 1.0}, "min": 0, "max": 10, "round": 1},
                                                                          "expected_salary_min": {"dist": "lognormal", "params": {"mu": 10.5, "sigma": 0.3}}}}]
    d["scenarios"].append(s2)
    cfg = parse_config(d)
    stats = []
    for sc in cfg.scenarios:
        ts = next(t for t in sc.tables if t.table == "candidate_profile")
        a, cols, _ = generate_batch(ctx, sc, ts, 0, 4000, 1)
        exp = np.array([float(x) for x in a["years_experience"] if x != NULL])
        sal = np.array([float(x) for x in a["expected_salary_min"] if x != NULL])
        stats.append((exp.mean(), exp.std(), np.log(sal).mean()))
    (m1, sd1, ls1), (m2, sd2, ls2) = stats
    assert abs(m1 - 8) < 0.2 and abs(sd1 - 2) < 0.2          # normal(8, 2) honoured
    assert abs(ls1 - 13.0) < 0.05 and abs(ls2 - 10.5) < 0.05  # lognormal mu honoured per scenario
    assert m2 < 2 and ls1 - ls2 > 2                           # scenario 2 is a clearly different population


def test_data_dictionary_enrichment(settings, tmp_path):
    import openpyxl
    from ingestion_engine.schema import load_data_dictionary
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Table", "Column", "Classification", "Description"])
    ws.append(["users", "email", "Personal", "login email"])
    ws.append(["users", "ghost", "Public", "not in db"])
    p = tmp_path / "dd.xlsx"
    wb.save(p)
    sch = Runner(settings).get_schema()
    assert load_data_dictionary(p, sch) == 1
    assert sch.tables["users"].columns["email"].classification == "Personal"
    assert load_data_dictionary(tmp_path / "missing.xlsx", sch) == 0
