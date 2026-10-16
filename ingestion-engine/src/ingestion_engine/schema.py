"""Live-DB schema introspection (information_schema / pg_catalog) + optional data dictionary."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import psycopg

ENGINE_TABLE_PREFIX = "talentra_ingest_"


@dataclass
class Column:
    name: str
    kind: str  # uuid,int,numeric,float,text,bool,date,timestamp,json,bytes,other
    sql_type: str
    nullable: bool
    default: str | None = None
    max_len: int | None = None
    precision: int | None = None
    scale: int | None = None
    enum_values: list[str] | None = None
    classification: str | None = None  # from data dictionary
    description: str | None = None

    @property
    def has_default(self) -> bool:
        return self.default is not None


@dataclass
class ForeignKey:
    columns: list[str]
    ref_table: str
    ref_columns: list[str]


@dataclass
class Table:
    name: str
    columns: dict[str, Column] = field(default_factory=dict)
    pk: list[str] = field(default_factory=list)
    fks: list[ForeignKey] = field(default_factory=list)
    uniques: list[list[str]] = field(default_factory=list)  # excluding PK
    checks: list[str] = field(default_factory=list)

    def fk_for(self, col: str) -> ForeignKey | None:
        for fk in self.fks:
            if col in fk.columns:
                return fk
        return None

    def parents(self) -> set[str]:
        return {fk.ref_table for fk in self.fks if fk.ref_table != self.name}


@dataclass
class Schema:
    tables: dict[str, Table]

    def order(self, names: list[str] | None = None) -> list[str]:
        """FK-aware topological load order (parents first). Cycles are broken deterministically."""
        want = set(names) if names is not None else set(self.tables)
        # include only dependencies inside `want`
        deps = {t: {p for p in self.tables[t].parents() if p in want} for t in want}
        out: list[str] = []
        remaining = dict(deps)
        while remaining:
            ready = sorted(t for t, d in remaining.items() if not (d - set(out)))
            if not ready:  # cycle: break by smallest dep count
                ready = [min(remaining, key=lambda t: (len(remaining[t] - set(out)), t))]
            for t in ready:
                out.append(t)
                remaining.pop(t)
        return out

    def to_dict(self) -> dict[str, Any]:
        return {
            n: {
                "pk": t.pk,
                "fks": [{"columns": f.columns, "ref_table": f.ref_table, "ref_columns": f.ref_columns} for f in t.fks],
                "uniques": t.uniques,
                "checks": t.checks,
                "columns": {
                    c.name: {"kind": c.kind, "type": c.sql_type, "nullable": c.nullable, "default": c.default,
                             "enum": c.enum_values, "classification": c.classification}
                    for c in t.columns.values()
                },
            }
            for n, t in self.tables.items()
        }


def _kind(data_type: str) -> str:
    d = data_type.lower()
    if d == "uuid":
        return "uuid"
    if d in ("bigint", "integer", "smallint", "int", "int2", "int4", "int8"):
        return "int"
    if d in ("numeric", "decimal"):
        return "numeric"
    if d in ("double precision", "real", "float", "float4", "float8"):
        return "float"
    if d in ("text", "character varying", "character", "string", "varchar", "char", "citext", "name"):
        return "text"
    if d == "boolean" or d == "bool":
        return "bool"
    if d == "date":
        return "date"
    if d.startswith("timestamp"):
        return "timestamp"
    if d in ("jsonb", "json"):
        return "json"
    if d == "bytea":
        return "bytes"
    return "other"


_QUOTED = re.compile(r"'((?:[^']|'')*)'")


def extract_enum_values(check_def: str) -> tuple[str, list[str]] | None:
    """`CHECK (col IN ('a','b'))` / `col = ANY (ARRAY['a','b'])` -> (col, values)."""
    m = re.search(r"\(?\s*\(?\"?(\w+)\"?(?:\)?::\w+)?\)?\s+(?:IN\s*\(|=\s*ANY\s*\(\s*(?:ARRAY)?\[?)", check_def, re.I)
    if not m:
        return None
    vals = [v.replace("''", "'") for v in _QUOTED.findall(check_def)]
    if not vals or " OR " in check_def.upper() and "IN" not in check_def.upper():
        return None
    return m.group(1), vals


_CMP = re.compile(r"\(?\"?(\w+)\"?\)?\s*(<=|>=|<|>)\s*\(?\"?(\w+)\"?\)?")


def ordered_pairs(table: Table) -> list[tuple[str, str]]:
    """Column pairs (lo, hi) such that a CHECK demands lo <= hi (or <). Used to repair generated rows."""
    pairs: list[tuple[str, str]] = []
    for chk in table.checks:
        for a, op, b in _CMP.findall(chk):
            if a in table.columns and b in table.columns and a != b:
                pairs.append((a, b) if op in ("<=", "<") else (b, a))
    return list(dict.fromkeys(pairs))


def introspect(conn: psycopg.Connection, schema: str = "public", include_engine_tables: bool = False) -> Schema:
    tables: dict[str, Table] = {}
    rows = conn.execute(
        """SELECT c.table_name, c.column_name, c.data_type, c.udt_name, c.is_nullable, c.column_default,
                  c.character_maximum_length, c.numeric_precision, c.numeric_scale, c.is_hidden,
                  c.generation_expression
           FROM information_schema.columns c
           JOIN information_schema.tables t ON t.table_schema=c.table_schema AND t.table_name=c.table_name
           WHERE c.table_schema=%s AND t.table_type='BASE TABLE'
           ORDER BY c.table_name, c.ordinal_position""", (schema,)).fetchall()
    for tn, cn, dt, udt, nullable, dflt, ml, prec, scale, hidden, gen in rows:
        if not include_engine_tables and tn.startswith(ENGINE_TABLE_PREFIX):
            continue
        if hidden == "YES" or gen or cn.startswith("crdb_internal"):
            continue
        t = tables.setdefault(tn, Table(tn))
        t.columns[cn] = Column(cn, _kind(dt), dt, nullable == "YES", dflt, ml, prec, scale)

    # pg_catalog lookups are split into small queries and joined in Python (a correlated-subquery version is ~100x slower on CRDB)
    rel = {oid: n for oid, n in conn.execute(
        "SELECT c.oid::int, c.relname FROM pg_class c JOIN pg_namespace n ON n.oid=c.relnamespace WHERE n.nspname=%s AND c.relkind='r'", (schema,)).fetchall()}
    attr = {(r, a): n for r, a, n in conn.execute(
        "SELECT a.attrelid::int, a.attnum::int, a.attname FROM pg_attribute a JOIN pg_class c ON c.oid=a.attrelid "
        "JOIN pg_namespace n ON n.oid=c.relnamespace WHERE n.nspname=%s AND a.attnum>0", (schema,)).fetchall()}
    defs = {oid: d for oid, d in conn.execute(
        "SELECT con.oid::int, pg_get_constraintdef(con.oid) FROM pg_constraint con JOIN pg_namespace n ON n.oid=con.connamespace "
        "WHERE n.nspname=%s AND con.contype='c'", (schema,)).fetchall()}
    cons = conn.execute(
        "SELECT con.oid::int, con.conrelid::int, con.contype::text, con.conkey, con.confrelid::int, con.confkey FROM pg_constraint con "
        "JOIN pg_namespace n ON n.oid=con.connamespace WHERE n.nspname=%s", (schema,)).fetchall()
    for coid, relid, ctype, ckey, frel, fkey in cons:
        tn = rel.get(relid)
        t = tables.get(tn) if tn else None
        if t is None:
            continue
        cols = [attr[(relid, k)] for k in (ckey or []) if (relid, k) in attr]
        cols = [c for c in cols if c in t.columns]
        if ctype == "p":
            t.pk = cols
        elif ctype == "u" and cols:
            t.uniques.append(cols)
        elif ctype == "f" and frel in rel:
            t.fks.append(ForeignKey(cols, rel[frel], [attr[(frel, k)] for k in (fkey or []) if (frel, k) in attr]))
        elif ctype == "c" and coid in defs:
            cdef = defs[coid]
            t.checks.append(cdef)
            ev = extract_enum_values(cdef)
            if ev and ev[0] in t.columns:
                t.columns[ev[0]].enum_values = ev[1]
    # unique indexes (CREATE UNIQUE INDEX) via information_schema.statistics
    idx = conn.execute(
        """SELECT table_name, index_name, non_unique, column_name, seq_in_index
           FROM information_schema.statistics WHERE table_schema=%s AND non_unique IN ('NO','0','f','false') AND (implicit IS NULL OR implicit='NO' OR implicit='f')
           ORDER BY table_name, index_name, seq_in_index""", (schema,)).fetchall()
    grouped: dict[tuple[str, str], list[str]] = {}
    for tn, iname, _nu, cn, _seq in idx:
        if iname == "primary" or iname.endswith("_pkey"):
            continue
        grouped.setdefault((tn, iname), []).append(cn)
    for (tn, _iname), cols in grouped.items():
        t = tables.get(tn)
        if not t:
            continue
        cols = [c for c in cols if c in t.columns]
        if cols and cols != t.pk and cols not in t.uniques:
            t.uniques.append(cols)
    return Schema(tables)


def load_data_dictionary(path: str | Path, schema: Schema) -> int:
    """Optionally enrich the introspected schema from ../docs/05_data_dictionary.xlsx.

    Looks for any sheet with header columns resembling table / column / classification / description.
    Returns number of columns enriched. Missing file or openpyxl -> 0 (feature is optional).
    """
    p = Path(path)
    if not p.exists():
        return 0
    try:
        import openpyxl  # type: ignore
    except ImportError:  # pragma: no cover
        return 0
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    n = 0
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        h = [str(x or "").strip().lower().replace(" ", "_") for x in header]

        def find(*names: str) -> int | None:
            for i, v in enumerate(h):
                if any(v == x or v.startswith(x) for x in names):
                    return i
            return None

        it, ic = find("table"), find("column", "field")
        icl, idesc = find("classification", "sensitivity"), find("description", "meaning")
        if it is None or ic is None:
            continue
        for r in rows:
            tn, cn = str(r[it] or "").strip(), str(r[ic] or "").strip()
            col = schema.tables.get(tn, Table(tn)).columns.get(cn)
            if col:
                if icl is not None and r[icl]:
                    col.classification = str(r[icl])
                if idesc is not None and r[idesc]:
                    col.description = str(r[idesc])
                n += 1
    return n
