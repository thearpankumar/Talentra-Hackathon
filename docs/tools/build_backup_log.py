#!/usr/bin/env python
"""Build docs/20_backup_restore_log.xlsx (B80 template).   usage: python docs/tools/build_backup_log.py"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
import xl
from dictlib import load_model
from openpyxl.styles import PatternFill, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "20_backup_restore_log.xlsx"))

# ------------------------------------------------------------------------------------------------ Plan
PLAN = [
 # section, item, laptop-scale (1M profiles), production target, notes / source
 ("Targets", "RTO (restore the service from backup after loss of the cluster or logical corruption)", 60, 30,
  "Minutes. Budget for the laptop test: detect + decide 5, prepare empty target cluster 10, RESTORE <= 30, validation 10, cutover 5 = 60. Production: restore of the hot set is parallelised across nodes; target 30 min assumes the production dataset is restored by a cluster at least as large as the source (measure and re-plan; extrapolate with the B74 growth model)."),
 ("Targets", "RPO (maximum acceptable data loss, minutes)", 15, 5,
  "Minutes. RPO of a backup-based recovery = age of the newest completed incremental at the failure instant; with revision_history the restore can target any time up to that point (PITR). Node/zone/region failure is NOT a backup case: the 3-replica cluster survives it with RPO 0 (survival goal ZONE or REGION)."),
 ("Targets", "Backup window", "15-minute incrementals, daily full", "5-minute incrementals, daily full, weekly full copied to a second bucket/account",
  "CockroachDB schedules: incrementals recur on the cron given in RECURRING, full on FULL BACKUP clause."),
 ("Targets", "Retention of backups", "7 days incrementals, 4 fulls (laptop: disk permitting)", "35 days PITR window (gc.ttlseconds not required for backup PITR), 13 monthly fulls (legal hold on audit/ledger evidence: 7 years via archive bucket with object lock)",
  "Align with data_retention_policy; erased subjects reappear in older backups: expire backups within 35 days or maintain an erasure re-apply list after a restore (DSR re-run job)."),
 ("Targets", "Restore test cadence", "Once before Gate 2 (B80), then each schema release", "Monthly full restore drill into the backup cluster; quarterly game-day with cutover rehearsal",
  "Record every run in sheet Log."),
 ("Scale", "Dataset (laptop seed)", "1,000,000 candidate profiles, 100,000 job requirements, 5,000,000 applications (docket section 6)", "10M profiles / 1M requirements / 100M applications (extrapolation target, not loaded)",
  "Physical size is measured at run time and written to the Log (GB). Planning assumption: replication factor 3 means the logical (backed-up) size is about one third of cluster disk use."),
 ("Method", "Backup type: full", "BACKUP INTO", "same",
  "Collection layout: each full creates a dated subdirectory under the collection URI; incrementals are appended to the latest full."),
 ("Method", "Backup type: incremental", "BACKUP INTO LATEST IN", "same",
  "Chain length bounded by the daily full."),
 ("Method", "Backup type: point-in-time (PITR)", "WITH revision_history on every full and incremental; RESTORE ... AS OF SYSTEM TIME", "same",
  "revision_history is on by default for scheduled backups (crdb_capabilities row 19) but the explicit option is kept for clarity."),
 ("Method", "Backup location", "S3-compatible MinIO container on the host or an S3 bucket (external connection backup_store). nodelocal://1/... for a quick laptop run.", "S3/GCS/Azure bucket in a different account/region than the cluster, object lock on the archive bucket, server-side KMS plus client KMS option",
  "crdb_capabilities: S3, GCS, Azure are the tested stores; MinIO-style S3 is 'untested' per the docs; nodelocal as a backup target is not described on the fetched pages (unverified): smoke-test it first."),
 ("Method", "Encryption of backups", "WITH encryption_passphrase (laptop) or KMS", "WITH kms = 'aws:///<key-arn>?AUTH=implicit&REGION=<r>' (or gcp:///...); separate key from the data-at-rest key; key policy allows backup_operator only",
  "Encrypted backups cannot use WITH EXPERIMENTAL COPY (Preview in v26.2, needs backup.compaction.threshold). Backup encryption options are from the official BACKUP docs, not from crdb_capabilities.md: verify on the pinned image. Application-encrypted columns (AES-GCM) stay ciphertext in every backup."),
 ("Method", "Licence note", "Multi-node cluster needs a Free/Trial/Enterprise licence (throttle: 5 open txns after 7 days without one). RESTORE no longer requires Enterprise. Whether BACKUP (and encrypted/KMS backups, schedules, revision_history) are licence-gated is UNVERIFIED on the fetched pages.", "Enterprise licence (Enterprise Free if the revenue test is met) applied after cockroach init; full-cluster backups carry the licence key",
  "Source: crdb_capabilities sections 1 row 19 and 3."),
 ("Method", "Separate backup cluster", "Second single-node (or 3-node) cluster started from the same image; used as restore target; no application connects to it", "Dedicated cluster in another region/account with its own CA; promoted only by the DR runbook",
  "infra/backup-db in the repository is the intended target; this sheet only documents the procedure."),
 ("Method", "What is backed up", "BACKUP INTO (cluster) for the DR drill; BACKUP DATABASE talentra for the dataset test", "cluster backup nightly + database/table-level on demand",
  "Roles/users live in system tables: restore them with RESTORE SYSTEM USERS (verify) or recreate with the Grants script; table privileges travel with the descriptors."),
 ("Security", "Who can run it", "backup_operator only (GRANT SYSTEM BACKUP / BACKUP ON DATABASE: verify privilege names)", "same + break-glass restore requires privileged_access_request",
  "See 13_rbac_matrix.xlsx, role backup_operator, E-26."),
 ("Security", "Backup content sensitivity", "Contains all ciphertext, hashes and Personal data in clear (columns not app-encrypted)", "same",
  "Backups are as sensitive as the database: encrypted, access logged, retention-limited, restore only on the backup cluster, no copy to laptops after the drill."),
]

# ------------------------------------------------------------------------------------------------ Procedure
PROC = [
 ("P-01", "Preflight", "Record cluster version, node count, licence state, dataset counts.", "SELECT version();  SHOW CLUSTER SETTING enterprise.license;  SELECT count(*) FROM candidate_profile; SELECT count(*) FROM job_requirement; SELECT count(*) FROM job_application;", "Counts match the seed plan (1M / 100k / 5M). Write them into Log notes.", "operator", ""),
 ("P-02", "Preflight", "Create the storage location (external connection).", "CREATE EXTERNAL CONNECTION backup_store AS 's3://talentra-backups/laptop?AUTH=implicit';   -- MinIO: add &AWS_ENDPOINT=http://minio:9000 (credentials per AUTH mode); quick test: 'nodelocal://1/backups'", "SHOW EXTERNAL CONNECTIONS lists backup_store.", "backup_operator", "Smoke-test write access with a tiny BACKUP of one lookup table."),
 ("P-03", "Preflight", "Create the schedule (production pattern; for the drill run P-04..P-05 manually).", "CREATE SCHEDULE talentra_backup FOR BACKUP INTO 'external://backup_store' WITH revision_history, encryption_passphrase = '<from vault>' RECURRING '*/15 * * * *' FULL BACKUP '@daily' WITH SCHEDULE OPTIONS first_run = 'now', on_previous_running = 'wait';", "SHOW SCHEDULES shows one full + one incremental schedule.", "backup_operator", "Replace the passphrase by kms = 'aws:///<arn>?AUTH=implicit&REGION=<r>' in production. Exact option names verified against the BACKUP / CREATE SCHEDULE docs on the pinned image."),
 ("P-04", "Backup", "Take a FULL backup of the seeded database and record start/end.", "BACKUP DATABASE talentra INTO 'external://backup_store' AS OF SYSTEM TIME '-10s' WITH revision_history, encryption_passphrase = '<from vault>';", "Job succeeds; note job id. Backup duration = finished - started.", "backup_operator", "Cluster form for DR: BACKUP INTO 'external://backup_store' AS OF SYSTEM TIME '-10s' WITH revision_history;"),
 ("P-05", "Backup", "Read the timing and size from the job record.", "SELECT job_id, job_type, status, created, started, finished, finished - started AS duration, fraction_completed FROM [SHOW JOBS] WHERE job_type = 'BACKUP' ORDER BY created DESC LIMIT 5;  SHOW BACKUPS IN 'external://backup_store';  SHOW BACKUP FROM LATEST IN 'external://backup_store' WITH encryption_passphrase = '<from vault>';", "Fill Log: backup start/end, size (GB) from SHOW BACKUP (size_bytes).", "operator", ""),
 ("P-06", "Backup", "Generate change traffic (insert 10k applications, 5k ledger rows) then take an INCREMENTAL backup; remember the wall-clock time T_fail (simulated failure instant).", "BACKUP DATABASE talentra INTO LATEST IN 'external://backup_store' AS OF SYSTEM TIME '-10s' WITH revision_history, encryption_passphrase = '<from vault>';", "Incremental finishes quickly; record end time; continue writing for a few minutes, then note T_fail = now().", "backup_operator", "RPO achieved = T_fail - (AS OF time of the newest backup)."),
 ("P-07", "Restore", "Prepare the target (separate backup cluster): empty cluster, same major version, same licence; connect as admin.", "SELECT version();  SHOW DATABASES;  -- talentra must not exist (or use RESTORE ... WITH new_db_name = 'talentra_restore')", "Target is empty.", "backup_operator", "Never restore into the production cluster for a drill."),
 ("P-08", "Restore", "RESTORE the database; record start/end.", "RESTORE DATABASE talentra FROM LATEST IN 'external://backup_store' WITH encryption_passphrase = '<from vault>';   -- PITR: ... FROM LATEST IN '...' AS OF SYSTEM TIME '<timestamp inside revision history>';", "Job succeeds; restore duration = finished - started. Single table drill: RESTORE TABLE talentra.public.wallet FROM LATEST IN '...' WITH into_db = 'talentra_restore';", "backup_operator", "RESTORE no longer requires an Enterprise licence (crdb_capabilities)."),
 ("P-09", "Restore", "Read restore timing.", "SELECT job_id, job_type, status, created, started, finished, finished - started AS duration FROM [SHOW JOBS] WHERE job_type = 'RESTORE' ORDER BY created DESC LIMIT 5;", "Fill Log: restore start/end, duration.", "operator", ""),
 ("P-10", "Restore", "Recreate roles and grants (login secrets are not part of table backups).", "RESTORE SYSTEM USERS FROM LATEST IN 'external://backup_store' WITH encryption_passphrase = '<from vault>';   -- verify availability; otherwise run the Grants script of 13_rbac_matrix.xlsx (GR-01..GR-05)", "SHOW ROLES lists the same roles as the source; SHOW GRANTS matches (AT-45).", "backup_operator", "Do not restore production service passwords/certificates to the drill cluster; issue new certificates."),
 ("P-11", "Validate", "Run every check of sheet Validation against source (AS OF SYSTEM TIME = backup time) and restored cluster; compare output.", "See sheet Validation (V-01..).", "All checks equal; hash-chain verifiers pass.", "operator", "Record result in Log: Pass / Pass with notes / Fail."),
 ("P-12", "Validate", "Application smoke test on the restored cluster: login, search, apply, wallet debit with the standard test accounts.", "Run the B65 query set subset (3 queries) and one write transaction in a rolled-back txn.", "Queries return expected counts; writes succeed.", "operator", "Counts toward RTO only if the target was ready (P-07)."),
 ("P-13", "Record", "Compute RTO met and RPO achieved (formulas in Log) and write the run in the Log; attach job ids and the Validation output.", "n/a", "RTO met = Yes when restore duration (+ validation duration when 'include validation' is ticked in Plan) <= RTO target.", "operator", ""),
 ("P-14", "Clean up", "Drop the restored database or destroy the drill cluster; revoke the temporary credentials; keep backups per retention.", "DROP DATABASE talentra CASCADE;  -- drill cluster only", "Target cluster empty; no drill data copied elsewhere.", "backup_operator", ""),
]

# ------------------------------------------------------------------------------------------------ Validation SQL
def validation_rows(model):
    T = {t["name"]: t for t in model["tables"]}
    rows = []
    n = 1

    def add(kind, desc, sql, expect, sev="High"):
        nonlocal n
        rows.append([f"V-{n:02d}", kind, desc, sql, expect, sev])
        n += 1

    add("Schema", "Table count identical", "SELECT count(*) FROM information_schema.tables WHERE table_schema = 'public' AND table_type = 'BASE TABLE';", f"= source ({len(T)} tables in the model)")
    add("Schema", "View count identical", "SELECT count(*) FROM information_schema.views WHERE table_schema = 'public';", "= source")
    add("Schema", "Constraint inventory identical", "SELECT constraint_type, count(*) FROM information_schema.table_constraints WHERE table_schema = 'public' GROUP BY 1 ORDER BY 1;", "= source for every type")
    add("Schema", "Index inventory identical", "SELECT count(*) FROM information_schema.statistics WHERE table_schema = 'public';", "= source", "Medium")
    add("Schema", "Foreign keys valid", "SELECT conname, conrelid::REGCLASS FROM pg_catalog.pg_constraint WHERE contype = 'f' AND NOT convalidated;", "0 rows")
    add("Security", "Roles identical", "SELECT username, member_of FROM [SHOW USERS] ORDER BY 1;", "= source minus drill-excluded logins")
    add("Security", "Grants identical on audit/ledger tables", "SHOW GRANTS ON TABLE audit_event, wallet_ledger_entry, profile_download, login_attempt;", "no UPDATE/DELETE for any grantee (13_rbac_matrix, AT-09)")
    add("Security", "No grants for PUBLIC on data tables", "SELECT table_name, privilege_type FROM information_schema.table_privileges WHERE grantee = 'public' AND table_schema = 'public';", "0 rows")
    big = [t for t in model["tables"] if not t["is_lookup"] and t.get("volume") in ("very_high", "high")]
    key = ["app_user", "candidate_profile", "job_requirement", "job_application", "organisation", "wallet", "wallet_ledger_entry", "payment_order", "profile_download",
           "audit_event", "audit_chain_head", "licence", "licence_seat", "device_session", "login_attempt", "notification", "notification_delivery", "user_block", "user_consent"]
    names = key + [t["name"] for t in big if t["name"] not in key]
    names = [x for x in names if x in T][:40]
    for nm in names:
        add("Row count + checksum", f"{nm}: count and id checksum equal to source at the backup time", f"SELECT count(*) AS n, sum(crc32c(id::STRING)) AS chk FROM {nm};   -- source: add AS OF SYSTEM TIME '<backup AS OF time>'", "n and chk equal", "High" if nm in key else "Medium")
    add("Row count", "Lookup tables: total rows over all lookups", "SELECT sum(n) FROM (" + " UNION ALL ".join(f"SELECT count(*) AS n FROM {t['name']}" for t in model["tables"] if t["is_lookup"] and t["module"] == "administration") + ") s;", "= source (extend to all 536 lookups with a generated script)", "Low")
    add("Integrity", "Wallet balances equal ledger", "SELECT * FROM vw_wallet_balance_drift LIMIT 10;", "0 rows (BR-SUB-08)")
    add("Integrity", "Subscription periods do not overlap", "SELECT * FROM vw_subscription_period_overlap LIMIT 10;", "0 rows (BR-SUB-03)")
    add("Integrity", "Invoice numbers gapless", "SELECT * FROM vw_invoice_number_gap LIMIT 10;", "0 rows (BR-SUB-12)")
    add("Integrity", "Audit hash chains verify", "-- run the chain verifier job: for each chain_no recompute row_hash = SHA-256(fields || prev_hash) and compare; compare newest chain_seq per chain_no with audit_chain_head", "no mismatch; zero security_event(audit_chain_break) created (BR-ADM-09)")
    add("Integrity", "Wallet ledger chain verifies", "-- nightly verifier for wallet_ledger_entry (row_hash chain per wallet) restricted to wallets touched in the last 24 h plus a 1% random sample", "no mismatch (BR-SUB-10)")
    add("Integrity", "One live subscription per licence", "SELECT licence_id, count(*) FROM subscription WHERE ended_at IS NULL AND deleted_at IS NULL GROUP BY 1 HAVING count(*) > 1;", "0 rows (BR-SUB-01)", "Medium")
    add("Integrity", "Seat/device limits hold", "SELECT device_slot_id, count(*) FROM device_session WHERE ended_at IS NULL GROUP BY 1 HAVING count(*) > 1;", "0 rows (BR-LIC-05)")
    add("Integrity", "Open recruiter mappings unique", "SELECT recruiter_app_user_id, organisation_id, count(*) FROM recruiter_organisation_mapping WHERE effective_to IS NULL GROUP BY 1, 2 HAVING count(*) > 1;", "0 rows (BR-ORG-01)", "Medium")
    add("Integrity", "Orphan check for deliberately FK-less references", "SELECT count(*) FROM audit_event a LEFT JOIN app_user u ON u.id = a.actor_app_user_id WHERE a.actor_app_user_id IS NOT NULL AND u.id IS NULL;", "0 (or equal to source)", "Medium")
    add("Integrity", "Blocked-party visibility still holds after restore", "SELECT count(*) FROM vw_candidate_visible_to_user WHERE viewer_app_user_id = '<blocked viewer>' AND candidate_app_user_id = '<blocker>';", "0 (AT-05)")
    add("Privacy", "Erased subjects are still erased", "SELECT count(*) FROM candidate_profile WHERE anonymised_at IS NOT NULL AND (contact_email IS NOT NULL OR contact_phone IS NOT NULL);", "0; if > 0 re-run the erasure job from candidate_erasure_request (backups may predate erasure)")
    add("Privacy", "Legal holds preserved", "SELECT count(*) FROM legal_hold WHERE released_at IS NULL;", "= source", "Medium")
    add("Time", "Timestamps in range (no clock-skew artefacts)", "SELECT max(created_at) FROM audit_event;  SELECT max(created_at) FROM wallet_ledger_entry;", "<= backup AS OF time", "Medium")
    add("Function", "Application smoke", "Run the three standard B65 queries and one wallet debit in a rolled-back txn", "same result shape and counts as source; debit succeeds", "High")
    return rows


def main():
    model = load_model()
    wb = xl.new_wb()
    # Plan
    ws = xl.add_sheet(wb, "Plan", ["section", "item", "laptop scale (1M profiles)", "production target", "notes / source"], PLAN, widths=[11, 44, 46, 46, 90], wrap_cols=("item", "laptop scale (1M profiles)", "production target", "notes / source"))
    # RTO/RPO named cells: row 2 = RTO, row 3 = RPO (columns C/D)
    base = ws.max_row + 2
    ws.cell(row=base, column=2, value="Summary of logged runs (reads sheet Log)")
    ws.cell(row=base, column=2).font = Font(bold=True)
    summ = [("Runs logged", "=COUNTA(Log!A2:A200)"),
            ("Runs with RTO met = Yes", '=COUNTIF(Log!P2:P200,"Yes")'),
            ("Runs with RTO met = No", '=COUNTIF(Log!P2:P200,"No")'),
            ("Longest restore (minutes)", '=IF(COUNT(Log!M2:M200)=0,"",MAX(Log!M2:M200))'),
            ("Longest backup (minutes)", '=IF(COUNT(Log!J2:J200)=0,"",MAX(Log!J2:J200))'),
            ("Worst RPO achieved (minutes)", '=IF(COUNT(Log!S2:S200)=0,"",MAX(Log!S2:S200))')]
    for i, (a, f) in enumerate(summ, start=1):
        ws.cell(row=base + i, column=2, value=a)
        ws.cell(row=base + i, column=3, value=f)
    ws.cell(row=base + 8, column=2, value="Include validation time in RTO? (Yes/No)")
    ws.cell(row=base + 8, column=3, value="Yes")
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    ws.add_data_validation(dv); dv.add(ws.cell(row=base + 8, column=3))
    inc_cell = f"Plan!$C${base + 8}"
    # Procedure
    xl.add_sheet(wb, "Procedure", ["step", "phase", "action", "SQL / command", "expected result", "role", "notes"], PROC, widths=[7, 11, 55, 110, 55, 16, 60], wrap_cols=("action", "SQL / command", "expected result", "notes"))
    # Log
    hdr = ["run_id", "date", "environment", "dataset rows (total)", "dataset size GB (logical, from SHOW BACKUP)", "backup type", "backup start", "backup end", "backup size GB",
           "backup duration (min)", "restore start", "restore end", "restore duration (min)", "validation duration (min)", "RTO target (min)", "RTO met", "simulated failure time",
           "restored-to (AS OF) time", "RPO achieved (min)", "RPO target (min)", "RPO met", "validation result", "operator", "notes"]
    rows = []
    for r in range(2, 62):
        rows.append([None, None, None, None, None, None, None, None, None,
                     f'=IF(AND(ISNUMBER(G{r}),ISNUMBER(H{r})),ROUND((H{r}-G{r})*1440,2),"")', None, None,
                     f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(L{r})),ROUND((L{r}-K{r})*1440,2),"")', None,
                     f'=IF(A{r}="","",Plan!$C$2)',
                     f'=IF(OR(M{r}="",O{r}=""),"",IF(M{r}+IF({inc_cell}="Yes",N{r},0)<=O{r},"Yes","No"))',
                     None, None,
                     f'=IF(AND(ISNUMBER(Q{r}),ISNUMBER(R{r})),ROUND((Q{r}-R{r})*1440,2),"")',
                     f'=IF(A{r}="","",Plan!$C$3)',
                     f'=IF(OR(S{r}="",T{r}=""),"",IF(S{r}<=T{r},"Yes","No"))',
                     None, None, None])
    ws = xl.add_sheet(wb, "Log", hdr, rows, widths=[10, 12, 18, 14, 16, 18, 18, 18, 11, 12, 18, 18, 12, 12, 11, 9, 18, 18, 11, 11, 9, 16, 16, 60], wrap_cols=("notes",),
                      validations={"environment": ["laptop-1M-seed", "laptop-sample", "staging", "prod-dr-drill"],
                                   "backup type": ["full", "incremental", "full+incremental", "PITR restore", "table restore", "cluster full"],
                                   "validation result": ["Pass", "Pass with notes", "Fail", "Not run"]})
    # formats
    for col, fmt in (("B", "yyyy-mm-dd"), ("G", "yyyy-mm-dd hh:mm:ss"), ("H", "yyyy-mm-dd hh:mm:ss"), ("K", "yyyy-mm-dd hh:mm:ss"), ("L", "yyyy-mm-dd hh:mm:ss"),
                     ("Q", "yyyy-mm-dd hh:mm:ss"), ("R", "yyyy-mm-dd hh:mm:ss"), ("D", "#,##0"), ("E", "0.00"), ("I", "0.00")):
        for r in range(2, 62):
            ws[f"{col}{r}"].number_format = fmt
    for col in ("J", "M", "N", "O", "S", "T"):
        for r in range(2, 62):
            ws[f"{col}{r}"].number_format = "0.00"
    # date / number validations
    for col, typ in (("B", "date"), ("G", "date"), ("H", "date"), ("K", "date"), ("L", "date"), ("Q", "date"), ("R", "date")):
        dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True, errorTitle="Date expected", error="Enter a valid date/time")
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}61")
    for col in ("D", "E", "I", "N"):
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, error="Enter a number >= 0")
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}61")
    # conditional formatting
    green, red, amber = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE"), PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE"), PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C")
    for col in ("P", "U"):
        ws.conditional_formatting.add(f"{col}2:{col}61", FormulaRule(formula=[f'{col}2="Yes"'], fill=green))
        ws.conditional_formatting.add(f"{col}2:{col}61", FormulaRule(formula=[f'{col}2="No"'], fill=red))
    ws.conditional_formatting.add("V2:V61", FormulaRule(formula=['V2="Pass"'], fill=green))
    ws.conditional_formatting.add("V2:V61", FormulaRule(formula=['V2="Pass with notes"'], fill=amber))
    ws.conditional_formatting.add("V2:V61", FormulaRule(formula=['V2="Fail"'], fill=red))
    # pre-filled readonly look for formula columns
    grey = PatternFill("solid", fgColor="F2F2F2")
    for col in ("J", "M", "O", "P", "S", "T", "U"):
        for r in range(2, 62):
            ws[f"{col}{r}"].fill = grey
    ws.freeze_panes = "B2"
    # Validation
    vr = validation_rows(model)
    xl.add_sheet(wb, "Validation", ["id", "type", "check", "SQL / action", "expected", "severity"], vr, widths=[7, 18, 60, 120, 50, 9], wrap_cols=("check", "SQL / action", "expected"))
    wb.save(OUT)
    print(f"wrote {OUT}: plan {len(PLAN)} rows, procedure {len(PROC)}, validation {len(vr)}")


if __name__ == "__main__":
    main()
