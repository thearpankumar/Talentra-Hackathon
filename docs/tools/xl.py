"""openpyxl helpers: consistently formatted sheets (header style, freeze panes, widths, filters, wrap)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
BAND = PatternFill("solid", fgColor="F2F5FA")
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def new_wb():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def add_sheet(wb, title, headers, rows, widths=None, wrap_cols=(), validations=None, status_col=None, max_width=70, freeze_cols=0):
    """headers: list[str]; rows: iterable of lists; validations: {col_name: [allowed values]}; status_col: header name to colour."""
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for r in rows:
        ws.append([_cell(v) for v in r])
    n = ws.max_row
    for c in range(1, len(headers) + 1):
        h = ws.cell(row=1, column=c)
        h.fill, h.font, h.border = HDR_FILL, HDR_FONT, BORDER
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    # widths
    for i, h in enumerate(headers, start=1):
        if widths and i - 1 < len(widths) and widths[i - 1]:
            w = widths[i - 1]
        else:
            mx = len(str(h))
            for row in ws.iter_rows(min_row=2, max_row=min(n, 400), min_col=i, max_col=i):
                v = row[0].value
                if v is not None:
                    mx = max(mx, min(len(str(v)), max_width))
            w = min(max(mx + 2, 10), max_width)
        ws.column_dimensions[get_column_letter(i)].width = w
    wrap_idx = {headers.index(c) + 1 for c in wrap_cols if c in headers}
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=(cell.column in wrap_idx))
    ws.freeze_panes = ws.cell(row=2, column=freeze_cols + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(n, 2)}"
    if validations:
        for col, allowed in validations.items():
            if col in headers and n >= 2:
                ci = get_column_letter(headers.index(col) + 1)
                dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{ci}2:{ci}{max(n, 2) + 500}")
    if status_col and status_col in headers:
        ci = get_column_letter(headers.index(status_col) + 1)
        rng = f"{ci}2:{ci}{max(n, 2) + 500}"
        for text, color in (("Done", "C6EFCE"), ("Complete", "C6EFCE"), ("Complies", "C6EFCE"), ("Yes", "C6EFCE"),
                            ("Partial", "FFEB9C"), ("In progress", "FFEB9C"), ("Deviation", "FFEB9C"),
                            ("Pending", "FFC7CE"), ("No", "FFC7CE"), ("Open", "FFC7CE"), ("Not started", "FFC7CE"),
                            ("Proposed", "DDEBF7"), ("Accepted", "C6EFCE")):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({ci}2,{len(text)})="{text}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))
    return ws


def _cell(v):
    if isinstance(v, (list, tuple, set)):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}={val}" for k, val in v.items())
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, str) and len(v) > 32000:
        return v[:32000]
    return v
