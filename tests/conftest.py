"""Shared fixtures.  Connection settings via environment (defaults = local insecure stack):

  TALENTRA_HOST=localhost TALENTRA_PORT=26267 TALENTRA_READ_PORT=26269 TALENTRA_BACKUP_PORT=26268
  TALENTRA_DB=talentra          database that holds the applied schema
  TALENTRA_USER=root            admin/superuser used for setup (insecure cluster: no password)
  TALENTRA_SSLMODE=disable
"""
import os
import uuid
from pathlib import Path

import openpyxl
import psycopg
import pytest

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
SCHEMA_DIR = ROOT / "infra" / "schema"
SYSTEM_ACTOR = "00000000-0000-0000-0000-000000000001"

HOST = os.environ.get("TALENTRA_HOST", "localhost")
PORT = int(os.environ.get("TALENTRA_PORT", "26267"))
READ_PORT = int(os.environ.get("TALENTRA_READ_PORT", "26269"))
BACKUP_PORT = int(os.environ.get("TALENTRA_BACKUP_PORT", "26268"))
DB = os.environ.get("TALENTRA_DB", "talentra")
USER = os.environ.get("TALENTRA_USER", "root")
SSL = os.environ.get("TALENTRA_SSLMODE", "disable")


def dsn(db=None, user=None, port=None):
    return f"postgresql://{user or USER}@{HOST}:{port or PORT}/{db or DB}?sslmode={SSL}"


def connect(db=None, user=None, port=None, role=None, autocommit=True):
    import time
    for i in range(30):  # HAProxy sometimes drops connections while a node is busy: reconnect
        try:
            c = psycopg.connect(dsn(db, user, port), autocommit=autocommit, connect_timeout=30)
            if role:
                c.execute(f"SET ROLE {role}")
            return c
        except psycopg.OperationalError:
            if i == 29:
                raise
            time.sleep(2)


def retry(fn, tries=10):
    """Run fn() retrying on 40001 (client-side retry loop required by SERIALIZABLE)."""
    import time
    for i in range(tries):
        try:
            return fn()
        except psycopg.errors.SerializationFailure:
            if i == tries - 1:
                raise
            time.sleep(0.05 * (i + 1))


@pytest.fixture(scope="session")
def admin():
    c = connect()
    yield c
    c.close()


@pytest.fixture()
def as_role():
    """as_role('app_service') -> autocommit connection that did SET ROLE <role> (root may assume any role)."""
    conns = []

    def make(role, autocommit=True):
        c = connect(role=role, autocommit=autocommit)
        conns.append(c)
        return c
    yield make
    for c in conns:
        try:
            c.close()
        except Exception:
            pass


def _load(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb:
        rows = list(ws.iter_rows(values_only=True))
        h = [str(x) if x is not None else "" for x in rows[0]]
        out[ws.title] = [dict(zip(h, ["" if v is None else v for v in r])) for r in rows[1:]]
    return out


@pytest.fixture(scope="session")
def dictionary():
    return _load(DOCS / "05_data_dictionary.xlsx")


@pytest.fixture(scope="session")
def rbac():
    return _load(DOCS / "13_rbac_matrix.xlsx")


@pytest.fixture(scope="session")
def rules():
    return _load(DOCS / "06_business_rules.xlsx")


@pytest.fixture(scope="session")
def catalog(admin):
    """Live catalog snapshot: columns, indexes, constraints, tables, views (public schema of TALENTRA_DB)."""
    cur = admin.cursor()
    cat = {}
    cat["tables"] = {r[0]: r[1] for r in cur.execute(
        "SELECT table_name, table_type FROM information_schema.tables WHERE table_schema='public'").fetchall()}
    cat["columns"] = {}
    for t, c, dt, nul, dflt, gen, pos in cur.execute(
            "SELECT table_name, column_name, data_type, is_nullable, column_default, generation_expression, ordinal_position "
            "FROM information_schema.columns WHERE table_schema='public' AND is_hidden='NO' ORDER BY table_name, ordinal_position").fetchall():
        cat["columns"].setdefault(t, []).append(dict(name=c, type=dt, nullable=nul == "YES", default=dflt, gen=gen, pos=pos))
    cat["indexes"] = {}
    for t, n, d in cur.execute("SELECT tablename, indexname, indexdef FROM pg_indexes WHERE schemaname='public'").fetchall():
        cat["indexes"].setdefault(t, {})[n] = d
    cat["constraints"] = {}
    for t, n, ty, d in cur.execute(
            "SELECT conrelid::regclass::string, conname, contype, pg_get_constraintdef(oid) FROM pg_constraint "
            "WHERE connamespace = 'public'::regnamespace AND conrelid <> 0").fetchall():
        cat["constraints"].setdefault(t, {})[n] = (ty, d)
    return cat


def new_id():
    return str(uuid.uuid4())
