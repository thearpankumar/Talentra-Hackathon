"""Helpers for the catalog-conformance suite (test_catalog_conformance.py, test_constraints_negative.py, ...).

* Dict(): the data dictionary sheets indexed by table (loaded once).
* norm_type / live_type: canonical type spelling for dictionary types and live catalog types.
* CheckSolver: uses the DATABASE AS ORACLE to (a) turn a row built by helpers_rowbuilder into a row that satisfies every CHECK
  of its table and (b) find, for ONE named CHECK, a mutation of the valid row that violates exactly that check.
  Candidate values are generated per column type; the check expressions are evaluated server side in one batched query
  (`SELECT (expr) IS FALSE FROM (VALUES ...)`), so no expression parsing is needed.
"""
import datetime
import decimal
import itertools
import re
import uuid
from functools import lru_cache
from pathlib import Path

import openpyxl
import psycopg
from psycopg.types.json import Jsonb

import helpers_rowbuilder as rb

DOCS = Path(__file__).resolve().parent.parent / "docs"


def yes(v):
    return str(v).strip().upper() == "YES"


class Dict:
    """Dictionary sheets + per-table indexes."""

    def __init__(self):
        wb = openpyxl.load_workbook(DOCS / "05_data_dictionary.xlsx", read_only=True, data_only=True)
        self.sheets = {}
        for ws in wb:
            rows = list(ws.iter_rows(values_only=True))
            h = [str(x) if x is not None else "" for x in rows[0]]
            self.sheets[ws.title] = [dict(zip(h, ["" if v is None else v for v in r])) for r in rows[1:]]
        S = self.sheets
        self.tables = {t["table"]: t for t in S["Tables"]}
        self.cols, self.idx, self.fks, self.cons, self.enums = {}, {}, {}, {}, {}
        for c in sorted(S["Columns"], key=lambda r: (r["table"], int(r["ordinal"]))):
            self.cols.setdefault(c["table"], []).append(c)
        for i in S["Indexes"]:
            self.idx.setdefault(i["table"], []).append(i)
        for f in S["ForeignKeys"]:
            self.fks.setdefault(f["table"], []).append(f)
        for c in S["Constraints"]:
            self.cons.setdefault(c["table"], []).append(c)
        for e in S["Enums"]:
            self.enums.setdefault(e["lookup_table"], []).append(e)
        self.views = {v["view_name"]: v for v in S["Views"]}

    def checks(self, table):
        return [c for c in self.cons.get(table, []) if c["constraint_type"] == "CHECK"]


@lru_cache(maxsize=1)
def get_dict():
    return Dict()


# ---------------------------------------------------------------------------------------------- types
def norm_type(t):
    """Dictionary type -> canonical string."""
    t = t.strip().upper().replace(" ", "")
    m = re.match(r"STRING\((\d+)\)$", t)
    if m:
        return f"STRING({m.group(1)})"
    m = re.match(r"DECIMAL\((\d+),(\d+)\)$", t)
    if m:
        return f"DECIMAL({m.group(1)},{m.group(2)})"
    return {"TIMESTAMPTZ": "TIMESTAMPTZ", "BOOL": "BOOL"}.get(t, t)


LIVE_BASE = {
    "uuid": "UUID", "boolean": "BOOL", "smallint": "INT2", "integer": "INT4", "bigint": "INT8", "text": "STRING",
    "timestamp with time zone": "TIMESTAMPTZ", "date": "DATE", "jsonb": "JSONB", "bytea": "BYTES", "inet": "INET",
    "time without time zone": "TIME", "interval": "INTERVAL", "numeric": "DECIMAL", "double precision": "FLOAT8",
    "real": "FLOAT4",
}


def live_type(r):
    """r = information_schema.columns row dict -> canonical type string."""
    dt = r["data_type"]
    if dt == "character varying" or dt == "character":
        n = r["character_maximum_length"]
        return f"STRING({n})" if n else "STRING"
    if dt == "numeric":
        if r["numeric_precision"] is None:
            return "DECIMAL"
        return f"DECIMAL({r['numeric_precision']},{r['numeric_scale']})"
    return LIVE_BASE.get(dt, dt.upper())


def pg_type(ft):
    """format_type() text -> canonical (CockroachDB prints STRING(n) as text(n+4))."""
    m = re.match(r"text\((\d+)\)$", ft)
    if m:
        return f"STRING({int(m.group(1)) - 4})"
    m = re.match(r"(?:character varying|varchar)\((\d+)\)$", ft)
    if m:
        return f"STRING({m.group(1)})"
    m = re.match(r"numeric\((\d+),(\d+)\)$", ft)
    if m:
        return f"DECIMAL({m.group(1)},{m.group(2)})"
    if ft == "numeric":
        return "DECIMAL"
    return LIVE_BASE.get(ft, ft.upper())


def cast_type(dtype):
    t = norm_type(dtype)
    if t.startswith("STRING"):
        return "STRING"
    if t.startswith("DECIMAL"):
        return "DECIMAL"
    return t


# ---------------------------------------------------------------------------------------------- candidates
def _str_cands(n, base):
    c = ["", " ", "x", "X", "a b", "ab", "AB", "ABC", "abc", "Abc", "ABCD", "12", "123", "0", "+", "+0123456", "+911234567890",
         "+1234567", "invalid", "UPPER", "Mixed@Example.COM", "a@b", "-", "a-b", "A_B", "a.b", "#", "1.0", "2999", "xx-YY-zzz",
         "25:00", "00:00", "*", "Z", "ftp://x", "http://x", "https://x.test", "en", "EN", "99", "-1", "000", "ZZ", "zz",
         "0123456789012345", "a" * 10, "A" * 3, "a" * n,
         "abc.com", "example.test", "a.b.co", "user@example.test", "USER@EXAMPLE.TEST", "+919876543210", "+12025550123", "#ff00aa",
         "https://example.test/x", "http://example.test", "a-b.example.com", "en-IN", "en-us", "Asia/Kolkata", "UTC", "192.168.1.1",
         "2025-01-01", "12:30", "abc123", "abc_123", "a1", "A1", "ab-cd", "INR", "USD", "in", "IN", "IND", "x.y", "a b c",
         "0.0.0.0", "00:11:22:33:44:55", "ABC-123", "abc-123", "v1", "1.0.0", "a/b", "/a", "a:b", "12345678", "+91", "+0",
         "https://", "www.example.test", "linkedin.com/in/x", "x" * 64]
    if base:
        c += [base.upper(), base.lower(), base + " ", " " + base, base[:-1], base + "!"]
    return [x for x in dict.fromkeys(c) if len(x) <= n]


def candidates(dtype, nullable, base, own_id, others):
    """Ordered candidate values for a column of dictionary type dtype."""
    t = norm_type(dtype)
    out = []
    if nullable:
        out.append(None)
    if t.startswith("STRING"):
        m = re.match(r"STRING\((\d+)\)", t)
        n = int(m.group(1)) if m else 400
        out += _str_cands(n, base if isinstance(base, str) else None)
    elif t in ("INT2", "INT4", "INT8"):
        mx = {"INT2": 32767, "INT4": 2147483647, "INT8": 2**40}[t]
        out += [0, -1, 1, 2, 3, 4, 5, 6, 7, 8, 10, 11, 12, 13, 24, 25, 31, 32, 50, 59, 60, 61, 99, 100, 101, 127, 128, 255, 256, 360,
                361, 365, 366, 367, 1000, 1001, 9999, 10000, -100, mx, mx - 1]
        if isinstance(base, int):
            out += [base - 1, base + 1, base * 2, base + 100]
        out = [x for x in out if x is None or -mx - 1 <= x <= mx]
    elif t.startswith("DECIMAL"):
        m = re.match(r"DECIMAL\((\d+),(\d+)\)", t)
        p, s = (int(m.group(1)), int(m.group(2))) if m else (18, 4)
        vals = ["0", "-1", "-0.01", "0.01", "0.5", "1", "1.01", "2", "5", "10", "50", "99.99", "100", "100.01", "101", "1000", "-100",
                "180.01", "-180.01", "90.01", "-90.01", "360.01", "1000000", "-0.0001", "0.0001", "3", "12", "24", "5.01", "4.01"]
        if isinstance(base, decimal.Decimal):
            vals += [str(base - 1), str(base + 1), str(base * 2), str(base + decimal.Decimal("0.01"))]
        lim = decimal.Decimal(10) ** (p - s)
        for v in vals:
            d = decimal.Decimal(v)
            if abs(d) < lim and d == d.quantize(decimal.Decimal(1).scaleb(-s)):
                out.append(d)
    elif t == "BOOL":
        out += [True, False]
    elif t == "DATE":
        b = base if isinstance(base, datetime.date) and not isinstance(base, datetime.datetime) else datetime.date.today()
        out += [b + datetime.timedelta(days=d) for d in (-1, 1, -7, 7, -30, 30, -365, 365, -3650, 3650)]
        out += [datetime.date(1900, 1, 1), datetime.date(1970, 1, 1), datetime.date(2100, 1, 1), datetime.date(3000, 1, 1)]
    elif t == "TIMESTAMPTZ":
        b = base if isinstance(base, datetime.datetime) else datetime.datetime.now(datetime.timezone.utc)
        out += [b + datetime.timedelta(**{k: v}) for k, v in
                (("seconds", -1), ("seconds", 1), ("minutes", -1), ("minutes", 1), ("hours", -1), ("hours", 1), ("days", -1), ("days", 1),
                 ("days", -30), ("days", 30), ("days", -365), ("days", 365), ("days", 3650), ("days", -3650))]
        utc = datetime.timezone.utc
        out += [datetime.datetime(1900, 1, 1, tzinfo=utc), datetime.datetime(2100, 1, 1, tzinfo=utc)]
    elif t == "JSONB":
        out += [Jsonb({}), Jsonb([]), Jsonb({"a": 1}), Jsonb("s"), Jsonb(1), Jsonb(None), Jsonb({"en": "x"}), Jsonb({"hi": "x"}),
                Jsonb([1, 2]), Jsonb({"x": {"y": 1}}), Jsonb(["a"]), Jsonb(True), Jsonb({"k": None})]
    elif t == "UUID":
        out += [uuid.UUID(str(own_id)) if own_id else uuid.uuid4(), uuid.uuid4()]
    elif t == "BYTES":
        out += [b"", b"x", b"\x00" * 8, b"\x00" * 16, b"\x00" * 32, b"\x00" * 64, b"x" * 3]
    elif t == "TIME":
        out += [datetime.time(0, 0), datetime.time(12, 0), datetime.time(23, 59, 59), datetime.time(23, 0)]
    elif t == "INTERVAL":
        out += [datetime.timedelta(0), datetime.timedelta(seconds=1), datetime.timedelta(days=1), datetime.timedelta(days=-1),
                datetime.timedelta(days=366)]
    elif t == "INET":
        out += ["10.0.0.1", "::1", "0.0.0.0"]
    # column-to-column copies of the same type
    for o in others:
        out.append(o)
    seen, res = set(), []
    for v in out:
        k = (type(v).__name__, repr(v) if not isinstance(v, Jsonb) else repr(v.obj))
        if k not in seen:
            seen.add(k)
            res.append(v)
    return res


_TOK = re.compile(r"[a-z_][a-z0-9_]*")
_STRLIT = re.compile(r"'(?:[^']|'')*'")


def refs_of(expr, colnames):
    e = _STRLIT.sub("''", expr)
    return [c for c in dict.fromkeys(_TOK.findall(e)) if c in colnames]


def insert_row(conn, table, row, colmap):
    names = list(row)
    ph = ", ".join("%s" + ("::INET" if norm_type(colmap[n]["type"]) == "INET" else "") for n in names)
    cur = conn.execute(f"INSERT INTO {table} ({', '.join(names)}) VALUES ({ph}) RETURNING *", [row[n] for n in names])
    r = cur.fetchone()
    return dict(zip([d.name for d in cur.description], r))


def build_valid_row(conn, table, _depth=0, _stack=(), **overrides):
    """Insert one row of `table` that satisfies all CHECKs (CheckSolver baseline) and return it (RETURNING *)."""
    for attempt in range(12):
        s = CheckSolver(conn, table)
        s.randomize = True
        explicit, _ = s.valid_baseline(populate=False, **overrides)
        try:
            return insert_row(conn, table, explicit, s.cols)
        except psycopg.errors.UniqueViolation:
            if attempt == 11:
                raise


_EXISTING = {}


class CheckSolver:
    """Per-table oracle + baseline row + violation search."""

    def __init__(self, conn, table, d=None):
        self.conn, self.table = conn, table
        self.d = d or get_dict()
        self.cols = {c["column"]: c for c in self.d.cols[table]}
        self.checks = {c["constraint_name"]: c["expression"] for c in self.d.checks(table)}
        self.computed = {n: re.sub(r"\s+(STORED|VIRTUAL)$", "", c["computed_expr"].strip()) for n, c in self.cols.items()
                         if (c.get("computed_expr") or "").strip()}
        names = set(self.cols)
        self.refs = {k: refs_of(e, names) for k, e in self.checks.items()}
        # physical columns that must be part of the oracle rows (check refs + sources of computed columns)
        need = set(itertools.chain.from_iterable(self.refs.values()))
        changed = True
        while changed:
            changed = False
            for n in list(need):
                if n in self.computed:
                    for s in refs_of(self.computed[n], names):
                        if s not in need:
                            need.add(s)
                            changed = True
        self.rel = [n for n in self.cols if n in need and n not in self.computed]
        self.mutable = {k: [c for c in dict.fromkeys(self._phys(v))] for k, v in self.refs.items()}
        self._defaults = {}
        self.randomize = False
        self.fkcols = {c.strip() for f in self.d.fks.get(table, []) for c in f["columns"].split(",")}
        self.fkref = {f["columns"].split(",")[0].strip(): f["ref_table"] for f in self.d.fks.get(table, [])}
        self.neigh = {}
        for k, rs in self.refs.items():
            ex = []
            for k2, rs2 in self.refs.items():
                if k2 != k and set(rs) & set(rs2):
                    ex += [x for x in self._phys(rs2) if x not in self.mutable[k] and x not in ex]
            self.neigh[k] = ex

    def _phys(self, refs):
        out = []
        for r in refs:
            if r in self.computed:
                out += [x for x in refs_of(self.computed[r], set(self.cols)) if x not in self.computed]
            else:
                out.append(r)
        return out

    # ------------------------------------------------------------------ base row
    def default_value(self, col):
        c = self.cols[col]
        d = str(c.get("default") or "").strip()
        if not d:
            return None
        if d not in self._defaults:
            self._defaults[d] = self.conn.execute(f"SELECT ({d})::{cast_type(c['type'])}").fetchone()[0]
        return self._defaults[d]

    def make_base(self, **overrides):
        """Plan a row via the row builder (parents are inserted), complete with defaults. Returns (explicit, full)."""
        old, rb.PARENT_BUILDER = rb.PARENT_BUILDER, build_valid_row
        try:
            plan = rb.build_row(self.conn, self.table, _plan=True, **overrides)
        finally:
            rb.PARENT_BUILDER = old
        for f in self.d.fks.get(self.table, []):
            fcols = [x.strip() for x in f["columns"].split(",")]
            rcols = [x.strip() for x in f["ref_columns"].split(",")]
            if len(fcols) < 2 or any(c not in plan for c in fcols) or f["ref_table"] == self.table:
                continue
            cond = " AND ".join(f"{rc} = %s" for rc in rcols)
            if self.conn.execute(f"SELECT 1 FROM {f['ref_table']} WHERE {cond} LIMIT 1", [plan[c] for c in fcols]).fetchone() is None:
                build_valid_row(self.conn, f["ref_table"], **{rc: plan[c] for c, rc in zip(fcols, rcols)})
        explicit = {"id": uuid.uuid4(), **plan} if "id" not in plan else dict(plan)
        full = dict(explicit)
        for n in self.cols:
            if n in full or n in self.computed:
                continue
            full[n] = self.default_value(n)
        return explicit, full

    # ------------------------------------------------------------------ oracle
    def evaluate(self, rows, names=None):
        """rows: list of full dicts. Returns list of dict check_name -> True when the check evaluates to FALSE."""
        names = list(names or self.checks)
        cols = self.rel
        ph = "(%s::INT, " + ", ".join(f"%s::{cast_type(self.cols[c]['type'])}" for c in cols) + ")"
        params = []
        for i, r in enumerate(rows):
            params.append(i)
            params += [r.get(c) for c in cols]
        comp = "".join(f", ({e}) AS {n}" for n, e in self.computed.items() if n in self._need_computed())
        flags = ", ".join(f"(({self.checks[n]}) IS FALSE)" for n in names)
        sql = (f"SELECT i, {flags} FROM (SELECT t.*{comp} FROM (VALUES {', '.join([ph] * len(rows))}) AS t(i, {', '.join(cols)})) u "
               "ORDER BY i")
        res = self.conn.execute(sql, params).fetchall()
        return [dict(zip(names, r[1:])) for r in res]

    def _need_computed(self):
        names = set()
        for k, e in self.checks.items():
            names |= {r for r in self.refs[k] if r in self.computed}
        return names

    # ------------------------------------------------------------------ search
    def _cands(self, col, row, mutated_cols, nonnull=False):
        c = self.cols[col]
        t = norm_type(c["type"])
        others = [row[o] for o in self.rel if o != col and o in row and row[o] is not None and self.cols[o]["type"] == c["type"]]
        nullable = yes(c["nullable"]) and col != "id" and not nonnull
        out = candidates(c["type"], nullable, row.get(col), row.get("id"), others)
        if self.randomize and col in self.fkcols and t == "UUID":
            ref = self.fkref.get(col)
            out = [x for x in out if x is None or (x == row.get("id") and ref == self.table)]   # no dangling FK values in baselines
            if ref and ref not in _EXISTING:
                r = self.conn.execute(f"SELECT id FROM {ref} LIMIT 1").fetchone()
                _EXISTING[ref] = r[0] if r else None
            if ref and _EXISTING.get(ref):
                out.insert(1 if nullable else 0, _EXISTING[ref])
        if self.randomize and t.startswith("STRING"):
            import random, string
            n = int(re.match(r"STRING\((\d+)\)", t).group(1)) if "(" in t else 40
            rnd = []
            for alphabet, ln in ((string.ascii_uppercase, 2), (string.ascii_uppercase, 3), (string.ascii_uppercase, 4),
                                 (string.ascii_lowercase, 3), (string.digits, 3), (string.digits, 4), (string.ascii_lowercase, 8)):
                if ln <= n:
                    rnd += ["".join(random.choice(alphabet) for _ in range(ln)) for _ in range(3)]
            rnd.append("".join(random.choice(string.ascii_lowercase) for _ in range(min(n, 12))))
            out = rnd + out
        return out

    def search(self, row, cols, accept, max_rows=4000, chunk=400, nonnull=False, must=None):
        """Find mutation of `cols` (1..3) in row such that accept(flags_dict) holds; returns dict of changes or None."""
        cols = [c for c in cols if c in self.cols and c not in self.computed]
        for size in (1, 2, 3):
            if size > len(cols):
                break
            for S in itertools.combinations(cols, size):
                if must and must not in S:
                    continue
                lists = [self._cands(c, row, S, nonnull) for c in S]
                if size == 3:
                    lists = [x[:9] for x in lists]
                combos = [x for x in itertools.product(*lists)]
                if len(combos) > max_rows:
                    combos = combos[:max_rows]
                for off in range(0, len(combos), chunk):
                    part = combos[off:off + chunk]
                    rows = [{**row, **dict(zip(S, v))} for v in part]
                    try:
                        flags = self.evaluate(rows)
                    except psycopg.Error:
                        continue
                    for v, f in zip(part, flags):
                        if accept(f):
                            return dict(zip(S, v))
        return None

    def populate(self, explicit, full):
        """Fill the nullable columns that CHECKs mention with non-NULL values (so single-column mutations violate exactly one
        CHECK instead of tripping pair checks such as (lat IS NULL) = (lon IS NULL))."""
        refd = []
        for k in self.checks:
            for c in self.mutable[k]:
                if c not in refd and c in self.cols and yes(self.cols[c]["nullable"]) and full.get(c) is None:
                    refd.append(c)
        fkmap = {f["columns"]: f for f in self.d.fks.get(self.table, []) if "," not in f["columns"]}
        for c in list(refd):
            f = fkmap.get(c)
            if f is None or f["ref_columns"] != "id":
                continue
            rt = f["ref_table"]
            try:
                if self.d.tables[rt]["kind"] == "lookup":
                    v = rb.lookup_id(self.conn, rt)
                elif rt == self.table:
                    continue
                else:
                    v = build_valid_row(self.conn, rt)["id"]
            except Exception:
                continue
            trial = {**full, c: v}
            if not any(self.evaluate([trial])[0].values()):
                full[c] = v
                explicit[c] = v
            refd.remove(c)
        rest = [c for c in refd if self.cols[c]["type"] and not (c in fkmap)]
        for c in rest:
            if full.get(c) is not None:
                continue
            fix = self.search(full, [c] + [x for x in rest if x != c][:2], lambda f: not any(f.values()), nonnull=True, max_rows=600, must=c)
            if fix:
                full.update(fix)
                explicit.update(fix)

    def valid_baseline(self, populate=True, **overrides):
        """-> (explicit_row, full_row) satisfying every CHECK, or raises AssertionError listing the unsatisfied ones."""
        explicit, full = self.make_base(**overrides)
        if not self.checks:
            return explicit, full
        for _ in range(10):
            fl = self.evaluate([full])[0]
            bad = [k for k, v in fl.items() if v]
            if not bad:
                if populate:
                    self.populate(explicit, full)
                for c in self.rel:   # pin every column the CHECKs look at (defaults such as now() must not drift)
                    explicit[c] = full.get(c)
                return explicit, full
            k = bad[0]
            fix = self.search(full, self.mutable[k] + self.neigh[k][:3], lambda f: not any(f.values()) or
                              (not f[k] and sum(f.values()) < len(bad)))
            if fix is None:
                raise AssertionError(f"{self.table}: cannot make baseline satisfy {k}")
            full.update(fix)
            explicit.update(fix)
        raise AssertionError(f"{self.table}: baseline did not converge")

    def violation(self, explicit, full, check):
        """Mutation (dict) violating ONLY `check`, or None."""
        others = [k for k in self.checks if k != check]
        return self.search(full, self.mutable[check], lambda f: f[check] and not any(f[o] for o in others))


class RConn:
    """Autocommit connection wrapper that reconnects and retries a statement when the (overloaded) local proxy drops the session."""

    def __init__(self, role=None, tries=40):
        self.role, self.tries, self._c = role, tries, None
        self._connect()

    def _connect(self):
        import conftest
        self._c = conftest.connect(role=self.role)
        self._c.execute("SET statement_timeout = '60s'")

    @property
    def info(self):
        return self._c.info

    @property
    def closed(self):
        return self._c.closed

    def close(self):
        self._c.close()

    def execute(self, sql, params=None):
        import time
        for i in range(self.tries):
            try:
                return self._c.execute(sql, params)
            except (psycopg.OperationalError, psycopg.errors.QueryCanceled):
                if i == self.tries - 1:
                    raise
                time.sleep(2)
                try:
                    self._connect()
                except psycopg.OperationalError:
                    pass

    def transaction(self):
        return self._c.transaction()
